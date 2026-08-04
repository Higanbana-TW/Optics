#!/usr/bin/env python3
"""Generate a Traditional-Chinese lens-distortion line calculator workbook."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path(__file__).with_name("lens_distortion_line_calculator.xlsx")


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "畸變線計算器"
    guide = wb.create_sheet("使用說明")

    navy = "17365D"
    blue = "D9EAF7"
    input_yellow = "FFF2CC"
    output_green = "E2F0D9"
    warning_red = "F4CCCC"
    white = "FFFFFF"
    gray = "E7E6E6"
    thin_gray = Side(style="thin", color="B7B7B7")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A14"
    ws.merge_cells("A1:N1")
    ws["A1"] = "Lens Distortion 歪斜線 4 點座標計算器"
    ws["A1"].font = Font(size=18, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:N2")
    ws["A2"] = "黃色儲存格為輸入；輸出以光軸為原點，亦可加上影像中心轉為絕對 pixel 座標。"
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    sections = {
        "A3:B3": "鏡頭／感測器輸入",
        "A13:C13": "Lens distortion table",
        "E13:F13": "輸入檢查",
        "A37:N37": "四點座標輸出",
    }
    for cells, text in sections.items():
        ws.merge_cells(cells)
        cell = ws[cells.split(":")[0]]
        cell.value = text
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center")

    input_rows = [
        (4, "Sensor pixel size", 3.45, "µm/pixel"),
        (5, "Lens 焦距", 4.0, "mm"),
        (6, "物體至 lens 距離", 1000.0, "mm"),
        (7, "影像中心 X", 960.0, "pixel"),
        (8, "影像中心 Y", 540.0, "pixel"),
        (9, "表格 FOV 定義", "全視場角", "全角或半角"),
    ]
    for row, label, value, unit in input_rows:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 3, unit)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=input_yellow)
        ws.cell(row, 2).border = border

    ws["A11"] = "有效表格列數"
    ws["B11"] = '=COUNT(A15:A34)'
    ws["A12"] = "表格狀態"
    ws["B12"] = (
        '=IF(B11<2,"錯誤：至少輸入 2 列",'
        'IF(SUMPRODUCT(--(A16:INDEX(A:A,14+B11)<=A15:INDEX(A:A,13+B11)))>0,'
        '"錯誤：FOV 必須遞增","OK"))'
    )

    line_inputs = [
        (4, "線段起點", -100.0, 50.0),
        (5, "線段終點", 100.0, 100.0),
    ]
    for cell, text in (("D3", "物方直線輸入"), ("E3", "X (mm)"), ("F3", "Y (mm)")):
        ws[cell] = text
        ws[cell].font = Font(bold=True, color=white)
        ws[cell].fill = PatternFill("solid", fgColor=navy)
        ws[cell].alignment = Alignment(horizontal="center")
    for row, label, x, y in line_inputs:
        ws.cell(row, 4, label)
        ws.cell(row, 5, x)
        ws.cell(row, 6, y)
        for col in (5, 6):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=input_yellow)
            ws.cell(row, col).border = border
    ws.merge_cells("D7:F10")
    ws["D7"] = (
        "座標假設：物體平面與 sensor 平行，光軸穿過物方 (0,0)。\n"
        "X 向右、物方 Y 向上；輸出影像 Y 為向下增加。\n"
        "4 點為線段 t = 0、1/3、2/3、1 的畸變後位置。"
    )
    ws["D7"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["D7"].fill = PatternFill("solid", fgColor=blue)

    ws["A14"] = "FOV (degree)"
    ws["B14"] = "Distortion (%)"
    ws["C14"] = "至下一列斜率"
    for cell in ws[14][0:3]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Example values are deliberately conspicuous and must be replaced by lens data.
    example_fov = [0, 20, 40, 60, 80]
    example_distortion = [0, -1, -3, -6, -10]
    for idx, row in enumerate(range(15, 35)):
        if idx < len(example_fov):
            ws.cell(row, 1, example_fov[idx])
            ws.cell(row, 2, example_distortion[idx])
        ws.cell(row, 3, f'=IF(OR(A{row}="",A{row + 1}=""),"",IFERROR((B{row + 1}-B{row})/(A{row + 1}-A{row}),""))')
        for col in (1, 2):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=input_yellow)
        for col in (1, 2, 3):
            ws.cell(row, col).border = border
    ws["A35"] = "注意"
    ws.merge_cells("B35:C35")
    ws["B35"] = "目前為示例值，請換成實際 lens table；Distortion 正值向外、負值向內。"
    ws["B35"].font = Font(color="9C0006", italic=True)
    ws["B35"].alignment = Alignment(wrap_text=True)

    checks = [
        (14, "Pixel size > 0", '=IF(B4>0,"OK","錯誤")'),
        (15, "焦距 > 0", '=IF(B5>0,"OK","錯誤")'),
        (16, "物距 > 0", '=IF(B6>0,"OK","錯誤")'),
        (17, "畸變表", "=B12"),
        (18, "FOV 範圍", '=IF(MAX(I39:I42)>MAX(A15:A34),"警告：超出表格，使用端點值","OK")'),
    ]
    for row, label, formula in checks:
        ws.cell(row, 5, label)
        ws.cell(row, 6, formula)
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=gray)
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=output_green)
        ws.cell(row, 5).border = ws.cell(row, 6).border = border

    headers = [
        "點",
        "t",
        "物方 X\n(mm)",
        "物方 Y\n(mm)",
        "理想像 X\n(mm)",
        "理想像 Y\n(mm)",
        "理想半徑\n(mm)",
        "半視場角\n(deg)",
        "查表 FOV\n(deg)",
        "內插 Dist.\n(%)",
        "畸變後 X\n(pixel,中心=0)",
        "畸變後 Y\n(pixel,中心=0)",
        "絕對 X\n(pixel)",
        "絕對 Y\n(pixel)",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(38, col, header)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[38].height = 42

    for point, row in enumerate(range(39, 43), 1):
        t = (point - 1) / 3
        ws.cell(row, 1, f"P{point}")
        ws.cell(row, 2, t)
        ws.cell(row, 3, f'=$E$4+B{row}*($E$5-$E$4)')
        ws.cell(row, 4, f'=$F$4+B{row}*($F$5-$F$4)')
        ws.cell(row, 5, f'=IFERROR($B$5*C{row}/$B$6,"")')
        ws.cell(row, 6, f'=IFERROR($B$5*D{row}/$B$6,"")')
        ws.cell(row, 7, f'=IFERROR(SQRT(E{row}^2+F{row}^2),"")')
        ws.cell(row, 8, f'=IFERROR(DEGREES(ATAN(G{row}/$B$5)),"")')
        ws.cell(row, 9, f'=IF($B$9="全視場角",2*H{row},H{row})')
        match_expr = f'MATCH(I{row},$A$15:INDEX($A$15:$A$34,$B$11),1)'
        ws.cell(
            row,
            10,
            f'=IF(OR($B$12<>"OK",I{row}=""),"",'
            f'IF(I{row}<=$A$15,$B$15,'
            f'IF(I{row}>=INDEX($A$15:$A$34,$B$11),INDEX($B$15:$B$34,$B$11),'
            f'INDEX($B$15:$B$34,{match_expr})+'
            f'(I{row}-INDEX($A$15:$A$34,{match_expr}))*'
            f'INDEX($C$15:$C$34,{match_expr}))))'
        )
        ws.cell(row, 11, f'=IFERROR(E{row}*(1+J{row}/100)*1000/$B$4,"")')
        # Pixel Y is positive downward, opposite to the object/sensor Y convention.
        ws.cell(row, 12, f'=IFERROR(-F{row}*(1+J{row}/100)*1000/$B$4,"")')
        ws.cell(row, 13, f'=IFERROR($B$7+K{row},"")')
        ws.cell(row, 14, f'=IFERROR($B$8+L{row},"")')
        for col in range(1, 15):
            ws.cell(row, col).border = border
            if col >= 3:
                ws.cell(row, col).fill = PatternFill("solid", fgColor=output_green)
            ws.cell(row, col).alignment = Alignment(horizontal="center")
            ws.cell(row, col).number_format = "0.000"

    ws["A44"] = "計算模型"
    ws.merge_cells("B44:N45")
    ws["B44"] = (
        "針孔投影：x=f·X/Z、y=f·Y/Z；徑向畸變：r_distorted=r_ideal·(1+Distortion%/100)。"
        "Distortion 依查表 FOV 線性內插；低於／高於表格時固定使用第一／最後一列，不外插。"
    )
    ws["B44"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B44"].fill = PatternFill("solid", fgColor=blue)

    # Plot absolute image coordinates. Reversing Y matches image-coordinate display.
    chart = ScatterChart()
    chart.title = "畸變後 4 點（影像 pixel 座標）"
    chart.style = 13
    chart.x_axis.title = "X (pixel)"
    chart.y_axis.title = "Y (pixel，向下為正)"
    chart.y_axis.scaling.orientation = "maxMin"
    chart.height = 10
    chart.width = 18
    x_values = Reference(ws, min_col=13, min_row=39, max_row=42)
    y_values = Reference(ws, min_col=14, min_row=39, max_row=42)
    series = Series(y_values, x_values, title="畸變後線段")
    series.graphicalProperties.line.solidFill = "C00000"
    series.graphicalProperties.line.width = 28575
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.graphicalProperties.solidFill = "C00000"
    chart.series.append(series)
    chart.legend = None
    ws.add_chart(chart, "A47")

    decimal_positive = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    decimal_positive.error = "請輸入大於 0 的數值"
    decimal_positive.errorTitle = "輸入錯誤"
    decimal_positive.showErrorMessage = True
    ws.add_data_validation(decimal_positive)
    decimal_positive.add("B4:B6")
    fov_choice = DataValidation(type="list", formula1='"全視場角,半視場角"')
    ws.add_data_validation(fov_choice)
    fov_choice.add(ws["B9"])

    ws.conditional_formatting.add(
        "F14:F18",
        FormulaRule(formula=['LEFT(F14,2)<>"OK"'], fill=PatternFill("solid", fgColor=warning_red)),
    )
    ws.conditional_formatting.add(
        "B12",
        CellIsRule(operator="notEqual", formula=['"OK"'], fill=PatternFill("solid", fgColor=warning_red)),
    )

    widths = {
        "A": 18,
        "B": 16,
        "C": 14,
        "D": 15,
        "E": 15,
        "F": 15,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 18,
        "L": 18,
        "M": 15,
        "N": 15,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = "A14:B34"

    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:F1")
    guide["A1"] = "使用說明與假設"
    guide["A1"].font = Font(size=18, bold=True, color=white)
    guide["A1"].fill = PatternFill("solid", fgColor=navy)
    guide["A1"].alignment = Alignment(horizontal="center")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 95
    instructions = [
        ("操作步驟", "1. 在黃色欄位輸入 pixel size、焦距、物距與影像中心。\\n"
         "2. 輸入物方直線的起點與終點 X/Y。\\n"
         "3. 以 lens datasheet 的 FOV/Distortion 取代示例表格，FOV 必須由小到大且不可留空列。\\n"
         "4. 由「四點座標輸出」讀取 P1～P4。"),
        ("FOV 定義", "若 datasheet 的 FOV 是左右（或對角線）兩側合計角度，選「全視場角」；"
         "若是從光軸到單側的 field angle，選「半視場角」。FOV 必須與物方 X/Y 所代表的同一像面方向相符。"),
        ("Distortion 定義", "本檔使用 Distortion(%)=(畸變像高/理想像高-1)×100%。"
         "正值是向外位移（通常稱 pincushion），負值是向內位移（通常稱 barrel）。"
         "若鏡頭資料表使用相反符號或 TV distortion，需先轉換，不能直接套用。"),
        ("座標系", "物方 X 向右、Y 向上，光軸穿過 (0,0)，物體平面與 sensor 平行。"
         "相對 pixel X 向右、Y 向下；絕對座標再加上影像中心。"),
        ("四點意義", "P1/P4 是起點/終點；P2/P3 位於線段的 1/3 與 2/3。"
         "徑向畸變後直線通常成曲線，四點只是三段折線近似；高畸變時應增加取樣點。"),
        ("適用限制", "模型只處理以光軸為中心、旋轉對稱的徑向 distortion。"
         "不包含 decentering/tangential distortion、sensor tilt、鏡頭姿態或不同深度的 3D 物體。"),
    ]
    for row, (title, body) in enumerate(instructions, 3):
        guide.cell(row, 1, title)
        guide.cell(row, 2, body)
        guide.cell(row, 1).font = Font(bold=True, color=white)
        guide.cell(row, 1).fill = PatternFill("solid", fgColor=navy)
        guide.cell(row, 2).fill = PatternFill("solid", fgColor=blue)
        guide.cell(row, 1).alignment = Alignment(vertical="top")
        guide.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row, 1).border = guide.cell(row, 2).border = border
        guide.row_dimensions[row].height = 62

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.active = 0
    return wb


if __name__ == "__main__":
    workbook = build_workbook()
    workbook.save(OUTPUT)
    print(f"Created {OUTPUT}")
