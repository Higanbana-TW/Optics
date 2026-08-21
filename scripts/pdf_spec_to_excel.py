#!/usr/bin/env python3
"""Convert the F139T-4 product specification PDF into a structured Excel workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "1F4E79"
NAVY_DARK = "0D2B4A"
WHITE = "FFFFFF"
LIGHT = "D6E3F0"
ZEBRA = "F4F8FB"
AMBER = "FFF2CC"
GREEN = "E2EFDA"
ORANGE = "FCE4D6"
RED_SOFT = "F8CBAD"
THIN = Border(
    left=Side(style="thin", color="B0B8C1"),
    right=Side(style="thin", color="B0B8C1"),
    top=Side(style="thin", color="B0B8C1"),
    bottom=Side(style="thin", color="B0B8C1"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(size=11, bold=False, color="000000", name="Calibri"):
    return Font(name=name, size=size, bold=bold, color=color)


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def apply_row(ws, row, values, *, fill_color=None, font_obj=None, align=WRAP, heights=None):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.border = THIN
        cell.alignment = align
        cell.font = font_obj or font()
        if fill_color:
            cell.fill = fill(fill_color)
    if heights:
        ws.row_dimensions[row].height = heights


def header_bar(ws, row, cols, title, fill_color=NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, title)
    cell.fill = fill(fill_color)
    cell.font = font(13, True, WHITE)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill(fill_color)
        ws.cell(row, c).border = THIN
    ws.row_dimensions[row].height = 24


def col_headers(ws, row, headers, fill_color=NAVY):
    apply_row(
        ws,
        row,
        headers,
        fill_color=fill_color,
        font_obj=font(11, True, WHITE),
        align=CENTER,
        heights=22,
    )


def kv_table(ws, start_row, rows, *, item_fill=LIGHT):
    r = start_row
    for i, (no, item, desc) in enumerate(rows):
        bg = WHITE if i % 2 == 0 else ZEBRA
        apply_row(ws, r, [no, item, desc], fill_color=bg, heights=max(36, 18 + str(desc).count("\n") * 14))
        ws.cell(r, 1).alignment = CENTER
        ws.cell(r, 1).font = font(11, True, NAVY)
        ws.cell(r, 2).fill = fill(item_fill if i % 2 == 0 else "EAF1F8")
        r += 1
    return r


def add_notes(ws, row, cols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = font(9, False, "666666")
    cell.alignment = WRAP
    ws.row_dimensions[row].height = 32 if "\n" in text else 18
    return row + 1


def build_workbook() -> Workbook:
    wb = Workbook()

    # ------------------------------------------------------------------ Cover
    ws = wb.active
    ws.title = "01 Cover"
    set_widths(ws, [22, 28, 42, 28, 22])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = "嘉兴中润光学科技股份有限公司  Jiaxing Zhongrun Optics Co., Ltd."
    c.font = font(14, True, WHITE)
    c.fill = fill(NAVY_DARK)
    c.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:E2")
    c = ws["A2"]
    c.value = "产品规格书  Product Specification"
    c.font = font(22, True, NAVY)
    c.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[2].height = 36

    ws.merge_cells("A3:E3")
    c = ws["A3"]
    c.value = "1/1.8\" 5X Zoom / Focus Lens  ·  IMX334  ·  8Mega HD"
    c.font = font(14, False, NAVY)
    c.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[3].height = 22

    header_bar(ws, 5, 5, "Document Information")
    col_headers(ws, 6, ["Field", "中文", "English / Value", "Unit", "Remark"])
    cover_rows = [
        ("Customer", "客户名称", "通用版 / General Edition", "", "Confidential"),
        ("Product Name", "产品名称", "1/1.8\" 5X Zoom/Focus Lens", "", ""),
        ("Product Model", "产品型号", "F139T-4", "", ""),
        ("Product P/N", "产品料号", "90.S08600.104", "", ""),
        ("Form No.", "表单编号", "Zmax/Q-SP-PE-01-37", "", "Rev1.0"),
        ("Version", "版本", "A", "", "2021-8-24 新规发行"),
        ("Company", "公司", "嘉兴中润光学科技股份有限公司", "", "浙江省嘉兴市陶泾路 188 号"),
        ("Tel", "电话", "0573-8222-9907", "", ""),
        ("Fax", "传真", "0573-8222-9909", "", ""),
        ("Drawn", "作成 Draw", "Hunter & Cole", "", "2021.08.24"),
        ("Reviewed", "审核 Review", "Skull", "", "2021.08.24"),
        ("Approved", "核准 Approval", "Peter / Owen", "", "2021.08.24"),
        ("Source PDF", "来源", "1/1.8 5X 4K IMX334 A 20210824", "24 pages", "Converted to Excel"),
    ]
    for i, row in enumerate(cover_rows):
        apply_row(ws, 7 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=22)
        ws.cell(7 + i, 1).font = font(11, True, NAVY)

    header_bar(ws, 21, 5, "Application  适用")
    ws.merge_cells("A22:E24")
    ws["A22"].value = (
        "本仕样书适用于 1/1.8″ CCD 或 CMOS 感光芯片，具备 5X 光学变焦以及自动对焦功能 8Mega 高清监控镜头。\n"
        "This specification is applicable to 8 Mega HD security lens with 1/1.8\" CCD or CMOS sensor and 5X Zoom & AF function."
    )
    ws["A22"].alignment = WRAP
    ws["A22"].font = font(11)
    ws.row_dimensions[22].height = 28
    ws.row_dimensions[23].height = 18
    ws.row_dimensions[24].height = 18

    header_bar(ws, 26, 5, "Workbook Map  工作表索引")
    col_headers(ws, 27, ["Sheet", "Section", "Content", "PDF pages", "Type"])
    index_rows = [
        ("01 Cover", "封面 / 适用", "Product identity, document control, application", "1–2", "Info"),
        ("02 Contents", "目录", "Original table of contents", "2", "Index"),
        ("03 Revision History", "1. 变更履历", "Revision record", "3", "Table"),
        ("04 Optical Spec", "3. 光学规格&性能", "Focal length, FOV, MTF, distortion", "4–5", "Table"),
        ("05 Zoom System", "4. 变焦群组", "Zoom motor, excitation map", "6–7", "Table"),
        ("06 Focus System", "5. 聚焦群组", "Focus motor, excitation map", "8–9", "Table"),
        ("07 IR Switch", "6. IR切换", "IR-cut / dummy glass spectral spec", "10–11", "Table"),
        ("08 Iris", "7. 光圈", "Iris motor parameters", "12", "Table"),
        ("09 Life Test", "8. 耐久测试", "Zoom / Focus / IR / Iris life cycles", "13–14", "Table"),
        ("10 Reliability", "9. 可靠度测试", "Storage, cycle, vibration, drop", "15–16", "Table"),
        ("11 FPC Spec", "10. FPC规格", "22-pin FPC pinout", "17", "Table"),
        ("12 Cam Lifting Map", "11. 驱动控制表", "Focus / Zoom travel vs steps", "18", "Table"),
        ("13 Mechanical Notes", "12–13. 外形/视点图", "Key dimensions from drawings", "19–21", "Notes"),
        ("14 View Point Data", "13. 2D视点图", "Max FOV and highest viewpoint data", "20–21", "Table"),
        ("15 Iris Fno Spec", "14. Iris&Fno规格", "Iris step 0–83 vs Fno / diameter / coil state", "22–23", "Numeric"),
        ("16 Appearance", "15. 外观检验规格", "MIL 60-40 appearance criteria", "24", "Spec"),
    ]
    for i, row in enumerate(index_rows):
        apply_row(ws, 28 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=20)
        ws.cell(28 + i, 1).font = font(11, True, NAVY)
        ws.cell(28 + i, 1).hyperlink = f"#'{row[0]}'!A1"
        ws.cell(28 + i, 1).font = Font(name="Calibri", size=11, bold=True, color="0563C1", underline="single")

    add_notes(
        ws,
        45,
        5,
        "Drawings on PDF pages 19–21 (mechanical outline, 2D viewpoint) are CAD graphics. "
        "Numeric callouts are captured on sheets 13 and 14; the original PDF remains the drawing authority.",
    )
    ws.freeze_panes = "A7"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddHeader.left.text = "Confidential"
    ws.oddHeader.right.text = "Zmax/Q-SP-PE-01-37 Rev1.0"
    ws.oddFooter.center.text = "F139T-4  |  90.S08600.104  |  Page &P of &N"

    # ------------------------------------------------------------------ Contents
    ws = wb.create_sheet("02 Contents")
    set_widths(ws, [10, 42, 42, 14])
    header_bar(ws, 1, 4, "目 录  Content")
    col_headers(ws, 2, ["No.", "中文", "English", "PDF"])
    contents = [
        ("1", "变更履历", "Revision History", "3"),
        ("2", "适用", "Application", "4"),
        ("3", "光学规格&性能", "Optical Specifications & Performance", "4"),
        ("4", "变焦群组", "Zoom Group / Zoom System", "6"),
        ("5", "聚焦群组", "Focus Group / Focus System", "8"),
        ("6", "IR切换", "IR Change Switch", "10"),
        ("7", "光圈", "Iris", "12"),
        ("8", "耐久测试", "Life Test", "13"),
        ("9", "可靠度测试", "Reliability Test", "15"),
        ("10", "FPC规格", "FPC Spec", "17"),
        ("11", "驱动控制表", "Cam Lifting Map", "18"),
        ("12", "机构外形图", "Mechanical Outside Drawing", "19"),
        ("13", "2D视点图", "2D View Point Drawing", "20"),
        ("14", "Iris&Fno规格", "Iris & Fno Spec", "22"),
        ("15", "外观检验规格", "Appearance Check Spec", "24"),
    ]
    for i, row in enumerate(contents):
        apply_row(ws, 3 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=20)
        ws.cell(3 + i, 1).alignment = CENTER
    add_notes(
        ws,
        19,
        4,
        "Original PDF TOC labels item 4 as Focus System and item 5 as Zoom System; body pages use Zoom Group then Focus Group. This workbook follows the body titles.",
    )
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Revision
    ws = wb.create_sheet("03 Revision History")
    set_widths(ws, [16, 12, 55, 16])
    header_bar(ws, 1, 4, "1. 变更履历  Revision History")
    col_headers(ws, 2, ["日期 Date", "版本 Ver.", "变更内容 Changes & Additions", "作成 Draw"])
    apply_row(
        ws,
        3,
        ["2021-8-24", "A", "新规发行  Preliminary version released", "Hunter"],
        fill_color=GREEN,
        heights=24,
        align=CENTER,
    )
    ws.cell(3, 3).alignment = WRAP

    # ------------------------------------------------------------------ Optical
    ws = wb.create_sheet("04 Optical Spec")
    set_widths(ws, [8, 28, 72])
    header_bar(ws, 1, 3, "3. 光学规格&性能  Optical Specifications & Performance")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"])
    optical = [
        (1, "焦距范围\nFocal Length", "10.5(±5%) ～ 47.0(±5%) mm"),
        (2, "F数\nF number", "F1.35(±5%) (Wide) ～ F1.55(±5%) (Tele)"),
        (3, "有效像圆径\nImage Circle", "Φ9.2 mm"),
        (4, "机械后焦\nMechanical Back Focus", "MBF = +0.76 mm (in air)"),
        (5, "摄影距离范围\nFocus Range", "Wide: 0.5 m ～ INF\nTele: 3.0 m ～ INF"),
        (6, "视场角\nField of View", "See FOV table (IMX334 Φ8.81 mm, 7.68 × 4.32)"),
        (7, "镜头机构外形\nOverall Size", "Φ40.0 mm × 83.9 mm"),
        (8, "镜头构成\nLens Structure", "4 群 15 枚 / 4 groups 15 elements"),
        (9, "相对照度\nRelative Illumination", ">30% (Image height = Φ9.2 mm)"),
        (10, "TV畸变\nTV Distortion", "-2.22% (Wide)  1.74% (Tele)\n(Image height = Φ8.81 mm, INF, SMIA)"),
        (11, "解像力\nResolution", "投影解像力 Projection Resolution. Measurement: 1 m (Wide) / 8 m (Tele); aperture = design value. Adjust focus at center then confirm periphery; balance for best contrast. Projection line pairs by internal standard chart — see table."),
        (12, "光晕&鬼影\nFlare & Ghost", "在实际使用中，无明显的鬼影或光斑。\nRemarkable ghost and flare should not be observed."),
        (13, "有害物质\nHazardous Substance", "依据 RoHS / Comply with RoHS"),
    ]
    kv_table(ws, 3, optical)
    header_bar(ws, 17, 3, "Field of View  IMX334  Φ8.81 mm (7.68 × 4.32)")
    col_headers(ws, 18, ["Axis", "Wide", "Tele"])
    for i, (axis, wide, tele) in enumerate([("D", "46.7°", "10.4°"), ("H", "41.0°", "9.1°"), ("V", "23.4°", "5.2°")]):
        apply_row(ws, 19 + i, [axis, wide, tele], fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=20)
        ws.cell(19 + i, 1).font = font(11, True, NAVY)
    header_bar(ws, 23, 3, "Projection Line Pairs  投影线对数")
    col_headers(ws, 24, ["Field", "Wide", "Tele"])
    for i, (field, wide, tele) in enumerate([("Center", "250 lp/mm", "250 lp/mm"), ("0.7 Field", "160 lp/mm", "160 lp/mm")]):
        apply_row(ws, 25 + i, [field, wide, tele], fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=20)
    add_notes(ws, 28, 3, "Measurement conditions: Distance 1 m (Wide) / 8 m (Tele). Aperture follows design value.")
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Zoom
    ws = wb.create_sheet("05 Zoom System")
    set_widths(ws, [8, 28, 72, 10, 10, 10, 10])
    header_bar(ws, 1, 7, "4. 变焦群组  Zoom Group")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description", "", "", "", ""])
    zoom = [
        (1, "动作方式\nOperation Method", "步进式马达驱动 / Stepping motor drive"),
        (2, "线圈电阻\nCoil Resistance", "55 Ω ± 5.5 Ω / phase (T=25℃)"),
        (3, "驱动电压\nOperation Voltage", "4.5 V ~ 5.0 V"),
        (4, "驱动转速\nDriving Speed", "800 pps"),
        (5, "激励方式\nExcitation Method", "2 相励磁 / 2 phase excitation"),
        (6, "步数角度\nStep Angle", "18°/step (2 phase excitation)"),
        (7, "传动螺杆\nScrew", "Pitch = 0.4 mm (when 2 phase excitation, 0.02 mm/step)"),
        (8, "背隙(水平)\nBacklash (horizontal)", "2 phase excitation, reverse direction: Zoom Group moves within 3 steps (including 3)."),
        (9, "工作温度\nOperating Temperature", "-30℃ ~ 70℃ (In a low temperature environment, warm up 30 minutes before start)."),
        (10, "噪音\nNoise", "60 dB(max), ambient 24 dB(max). Probe–motor 10 cm, 5 V, 800 pps."),
        (11, "光耦\nOptocouple", "Type: ROHM RPI-222\nVH = 2.3 V over; VL = 0.7 V under"),
        (12, "动作方向\nOperation Direction", "See 2-phase excitation map. CW: Tele→Wide; CCW: Wide→Tele"),
    ]
    kv_table(ws, 3, zoom)
    header_bar(ws, 16, 7, "2相励磁  2-phase excitation  ·  FPC & Motor Terminal")
    col_headers(ws, 17, ["Pin / Coil", "1", "2", "3", "4", "Direction", ""])
    for i, row in enumerate(
        [
            ["12  B-", "H", "H", "L", "L", "CW: Tele → Wide", ""],
            ["13  A+", "H", "L", "L", "H", "", ""],
            ["14  B+", "L", "L", "H", "H", "CCW: Wide → Tele", ""],
            ["15  A-", "L", "H", "H", "L", "", ""],
        ]
    ):
        apply_row(ws, 18 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=22)
        for col in range(2, 6):
            if ws.cell(18 + i, col).value == "H":
                ws.cell(18 + i, col).fill = fill(GREEN)
            else:
                ws.cell(18 + i, col).fill = fill(AMBER)
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Focus
    ws = wb.create_sheet("06 Focus System")
    set_widths(ws, [8, 28, 72, 8, 8, 8, 8, 8, 8, 8, 8])
    header_bar(ws, 1, 11, "5. 聚焦群组  Focus Group")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"] + [""] * 8)
    focus = [
        (1, "动作方式\nOperation Method", "步进式马达驱动 / Stepping motor drive"),
        (2, "线圈电阻\nCoil Resistance", "60 Ω ± 6 Ω / phase (T=25℃)"),
        (3, "驱动电压\nOperation Voltage", "4.5 V ~ 5.0 V"),
        (4, "驱动转速\nDriving Speed", "800 pps"),
        (5, "励磁方式\nExcitation Method", "1-2 相励磁 / 1-2 phase excitation"),
        (6, "步数角度\nStep Angle", "9°/step (1-2 phase excitation)"),
        (7, "传动螺杆\nScrew", "Pitch = 0.4 mm (when 1-2 phase excitation, 0.01 mm/step)"),
        (8, "背隙(水平)\nBacklash (horizontal)", "1-2 phase excitation, reverse direction: Focus Group moves within 6 steps (including 6)."),
        (9, "工作温度\nOperating Temperature", "-30℃ ~ 70℃ (warm up 30 minutes at low temperature before start)."),
        (10, "噪音\nNoise", "60 dB(max), ambient 24 dB(max). Probe–motor 10 cm, 5 V, 800 pps."),
        (11, "光耦\nOptocouple", "Type: ROHM RPI-222\nVH = 2.3 V over; VL = 0.7 V under"),
        (12, "动作方向\nOperation Direction", "See 1-2 phase excitation map. CW: Near→Far; CCW: Far→Near"),
    ]
    kv_table(ws, 3, focus)
    header_bar(ws, 16, 11, "1-2相励磁  1-2 phase excitation  ·  FPC & Motor Terminal")
    col_headers(ws, 17, ["Pin / Coil", "1", "2", "3", "4", "5", "6", "7", "8", "Direction", ""])
    for i, row in enumerate(
        [
            ["08  B-", "L", "L", "H", "H", "H", "L", "L", "L", "CW: Near → Far", ""],
            ["09  A+", "H", "H", "H", "L", "L", "L", "L", "L", "", ""],
            ["10  B+", "H", "L", "L", "L", "L", "L", "H", "H", "CCW: Far → Near", ""],
            ["11  A-", "L", "L", "L", "L", "H", "H", "H", "L", "", ""],
        ]
    ):
        apply_row(ws, 18 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=22)
        for col in range(2, 10):
            if ws.cell(18 + i, col).value == "H":
                ws.cell(18 + i, col).fill = fill(GREEN)
            else:
                ws.cell(18 + i, col).fill = fill(AMBER)
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ IR
    ws = wb.create_sheet("07 IR Switch")
    set_widths(ws, [8, 32, 72])
    header_bar(ws, 1, 3, "6. IR 切换  IR Change Switch")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"])
    ir = [
        (1, "电阻\nResistance", "20 Ω ± 3 Ω (T=20℃)"),
        (2, "作动电压\nOperation Voltage", "4.5 V ≤ |U| ≤ 5.0 V"),
        (3, "作动时间\nOperation Time", "200 ms ~ 500 ms"),
        (
            4,
            "作动方向\nOperation Direction",
            "PIN2 +, PIN1 − : Night mode → Day mode\n"
            "PIN2 −, PIN1 + : Day mode → Night mode\n"
            "Note: Need refocus after switching IR filter.",
        ),
        (5, "IR 规格\nIR Spec", "Two filters, each t = 0.3 ± 0.05 mm. Spectral tables below."),
    ]
    kv_table(ws, 3, ir)
    header_bar(ws, 9, 3, "IR Filter (IR-cut / Day)  t = 0.3 ± 0.05 mm")
    col_headers(ws, 10, ["项目 Item", "Wavelength", "规格 Spec"])
    ir_cut = [
        ("透过率 transmissivity", "350–395 nm", "Tave < 3%"),
        ("透过率 transmissivity", "415 ± 10 nm", "T = 50%"),
        ("透过率 transmissivity", "430–565 nm", "Tave > 93%"),
        ("透过率 transmissivity", "650 ± 10 nm", "T = 50%"),
        ("透过率 transmissivity", "700–725 nm", "Tave < 5%"),
        ("透过率 transmissivity", "725–1100 nm", "Tave < 0.5%"),
    ]
    for i, row in enumerate(ir_cut):
        apply_row(ws, 11 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=20)
    header_bar(ws, 18, 3, "Dummy glass / Night  t = 0.3 ± 0.05 mm")
    col_headers(ws, 19, ["项目 Item", "Wavelength", "规格 Spec"])
    dummy = [
        ("透过率 transmissivity", "430–680 nm", "Tave ≥ 97%，Tmin ≥ 95%"),
        ("透过率 transmissivity", "680–1100 nm", "Tave ≥ 96%，Tmin ≥ 94%"),
    ]
    for i, row in enumerate(dummy):
        apply_row(ws, 20 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=20)
    add_notes(ws, 23, 3, "Page 11 of the PDF continues the IR spec with the second (broadband) filter table; labels Day/Night follow typical dual-filter usage of the tabulated spectra.")
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Iris
    ws = wb.create_sheet("08 Iris")
    set_widths(ws, [8, 32, 72])
    header_bar(ws, 1, 3, "7. 光圈  Iris")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"])
    iris = [
        (1, "叶片数量\nBlades Number", "2 枚叶片 / 2 slices"),
        (2, "驱动方式\nDrive Mode", "PM 型步进电机 / PM Stepping Motor"),
        (3, "物理&光学最大口径\nPhysical & Optical Maximum Diameter", "Physical Φ17.4 mm\nOptical Φ16.4 mm"),
        (4, "线圈电阻\nCoil Resistance", "30 Ω ± 5 Ω (T=25℃)"),
        (5, "驱动电压\nOperation Voltage", "2.7 V ~ 5.0 V"),
        (6, "驱动转速\nDriving Speed", "200 pps"),
        (7, "励磁方式\nExcitation Method", "2-2 相励磁 / 2-2 phase excitation"),
        (8, "步数角度\nStep Angle", "0.845°/step (2-2 phase excitation)"),
        (9, "噪音\nNoise", "74 dB(max), ambient 24 dB(max). Probe–motor 25 cm, 2.6 V, 160 pps."),
        (10, "Iris&Fno 规格\nIris & Fno Spec", "详见第 14 项 / Details on sheet 15 Iris Fno Spec"),
    ]
    kv_table(ws, 3, iris)
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Life
    ws = wb.create_sheet("09 Life Test")
    set_widths(ws, [8, 32, 28, 22, 55])
    header_bar(ws, 1, 5, "8. 耐久测试  Life Test")
    col_headers(ws, 2, ["NO.", "项目 Item", "Cycles", "Environment", "Pass criterion / Notes"])
    life = [
        (
            1,
            "变焦群组耐久测试\nZoom Group Life Test",
            "500,000",
            "25±5℃ / 55±20% RH",
            "Parameters per Zoom Group (item 4). After test, lens has no performance issues.",
        ),
        (
            2,
            "聚焦群组耐久测试\nFocus Group Life Test",
            "500,000",
            "25±5℃ / 55±20% RH",
            "Parameters per Focus Group (item 5). After test, lens has no performance issues.",
        ),
        (
            3,
            "IR切换耐久测试\nIR Change Life Test",
            "50,000",
            "Normal temp & humidity",
            "Close → Open → Close, 4.0–6.0 s/cycle. After test, IR-cut has no performance issues.",
        ),
        (
            4,
            "Iris耐久测试\nIris Life Test",
            "300,000",
            "Normal temp & humidity",
            "Open → Close → Open, 4–8 s/cycle. After test, Iris has no performance issues.",
        ),
    ]
    for i, row in enumerate(life):
        apply_row(ws, 3 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=48)
        ws.cell(3 + i, 1).alignment = CENTER
        ws.cell(3 + i, 3).alignment = CENTER
        ws.cell(3 + i, 3).font = font(12, True, NAVY)
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Reliability
    ws = wb.create_sheet("10 Reliability")
    set_widths(ws, [8, 32, 78])
    header_bar(ws, 1, 3, "9. 可靠度测试  Reliability Test")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"])
    rel = [
        (
            1,
            "低温保存\nLow Temperature Storage",
            "-30℃ ± 3℃, 48 h. Recover 12–24 h to room temperature/humidity. No performance issues.",
        ),
        (
            2,
            "高温保存\nHigh Temperature Storage",
            "70℃ ± 3℃, humidity 50% ± 5%, 48 h. Recover 12–24 h. No performance issues.",
        ),
        (
            3,
            "高温高湿保存\nHigh Temperature & Humidity Storage",
            "70℃ ± 3℃, humidity 85% ± 5%, 48 h. Recover 12–24 h. No performance issues.",
        ),
        (
            4,
            "温度循环测试\nTemperature Cycle Test",
            "Condition 1: +60℃, 30 min, humidity 50%.\n"
            "Condition 2: −20℃, 30 min.\n"
            "Average temperature change rate 10℃/min. Repeat 30 cycles. Recover 12–24 h. No performance issues.",
        ),
        (
            5,
            "震动测试\nVibration Test",
            "Lens with package. PSD profile in table. Acceleration 1.087 Grms, 3 axes: Z 60 min, X and Y 30 min each. After test, no performance issues.",
        ),
        (
            6,
            "落下测试\nDrop Test",
            "Lens with package. Dropped from 1.0 m. 6 surfaces + 3 edges + 1 corner, once each, 10 times total. After test, no performance issues.",
        ),
    ]
    kv_table(ws, 3, rel)
    header_bar(ws, 10, 3, "Vibration PSD  震动谱密度")
    col_headers(ws, 11, ["频率 Frequency (Hz)", "加速度谱密度 PSD (g²/Hz)", ""])
    psd = [(5.0, 0.01), (100.0, 0.01), (135.0, 0.00148), (200.0, 0.001), (300.0, 0.00001)]
    for i, (hz, g) in enumerate(psd):
        apply_row(ws, 12 + i, [hz, g, ""], fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=20)
        ws.cell(12 + i, 1).number_format = "0.0"
        ws.cell(12 + i, 2).number_format = "0.00000"
    add_notes(ws, 18, 3, "Total RMS acceleration 1.087 Grms.")
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ FPC
    ws = wb.create_sheet("11 FPC Spec")
    set_widths(ws, [10, 28, 18, 40])
    header_bar(ws, 1, 4, "10. FPC 规格  FPC Spec")
    col_headers(ws, 2, ["No.", "FUNCTION", "ITEM", "Note"])
    fpc = [
        (1, "IR+", "IR", "金手指方向朝上 / Gold finger upward"),
        (2, "IR-", "IR", ""),
        (3, "PI ANODE", "FOCUS PI", ""),
        (4, "PI CATHODE&EMITTER", "FOCUS PI", ""),
        (5, "None", "", ""),
        (6, "None", "", ""),
        (7, "PI COLLECTOR(FOCUS)", "FOCUS PI", ""),
        (8, "B-(FOCUS)", "FOCUS Motor", ""),
        (9, "A+(FOCUS)", "FOCUS Motor", ""),
        (10, "B+(FOCUS)", "FOCUS Motor", ""),
        (11, "A-(FOCUS)", "FOCUS Motor", ""),
        (12, "B-(ZOOM)", "ZOOM Motor", ""),
        (13, "A+(ZOOM)", "ZOOM Motor", ""),
        (14, "B+(ZOOM)", "ZOOM Motor", ""),
        (15, "A-(ZOOM)", "ZOOM Motor", ""),
        (16, "PI ANODE", "ZOOM PI", ""),
        (17, "PI CATHODE&EMITTER", "ZOOM PI", ""),
        (18, "PI COLLECTOR(ZOOM)", "ZOOM PI", ""),
        (19, "A+ (IRIS)", "IRIS", ""),
        (20, "B+ (IRIS)", "IRIS", ""),
        (21, "A- (IRIS)", "IRIS", ""),
        (22, "B- (IRIS)", "IRIS", ""),
    ]
    group_fill = {
        "IR": ORANGE,
        "FOCUS PI": LIGHT,
        "FOCUS Motor": "C6EFCE",
        "ZOOM Motor": "BDD7EE",
        "ZOOM PI": "DDEBF7",
        "IRIS": AMBER,
        "": ZEBRA,
    }
    for i, (no, fn, item, note) in enumerate(fpc):
        apply_row(ws, 3 + i, [no, fn, item, note], fill_color=group_fill.get(item, WHITE), heights=20)
        ws.cell(3 + i, 1).alignment = CENTER
        ws.cell(3 + i, 1).font = font(11, True, NAVY)
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ Cam map
    ws = wb.create_sheet("12 Cam Lifting Map")
    set_widths(ws, [12, 55, 18, 14, 18])
    header_bar(ws, 1, 5, "11. 驱动控制表  Cam Lifting Map")
    apply_row(
        ws,
        2,
        ["① Zoom Step: 0.02 mm/Step", "② Focus Step: 0.01 mm/Step", "", "", ""],
        fill_color=AMBER,
        font_obj=font(11, True, NAVY),
        heights=22,
    )
    col_headers(ws, 3, ["项目 Item", "说明 Description", "尺寸 Dimension (mm)", "步数 Step", "Check (mm/step)"])
    cam = [
        ("F1", "FOCUS: TELE INF(VIS) ～ Focus Sensor", 1.57, 157, 0.01),
        ("F2", "FOCUS: TELE INF(VIS) ～ Focus (NEAR) 3 m IR", 4.34, 434, 0.01),
        ("F3", "FOCUS: Near 端机构余量 (Near end mechanism margin)", 0.78, 78, 0.01),
        ("F4", "FOCUS: Far 端机构余量 (Far end mechanism margin)", 0.97, 97, 0.01),
        ("F5", "FOCUS: TELE INF(VIS) ～ WIDE INF(VIS)", 1.29, 129, 0.01),
        ("Z1", "ZOOM: TELE ～ Focus (NEAR) 3 m IR", 6.12, 306, 0.02),
        ("Z2", "ZOOM: TELE 端 ～ Zoom sensor", 5.56, 278, 0.02),
        ("Z3", "ZOOM: TELE — WIDE", 19.42, 971, 0.02),
        ("Z4", "ZOOM: TELE 端机构余量 (TELE end mechanism margin)", 0.94, 47, 0.02),
        ("Z5", "ZOOM: WIDE 端机构余量 (WIDE end mechanism margin)", 0.60, 30, 0.02),
    ]
    for i, (item, desc, dim, step, pitch) in enumerate(cam):
        r = 4 + i
        apply_row(ws, r, [item, desc, dim, step, f"=C{r}/D{r}"], fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=22)
        ws.cell(r, 1).font = font(11, True, NAVY)
        ws.cell(r, 1).alignment = CENTER
        ws.cell(r, 3).number_format = "0.00"
        ws.cell(r, 4).number_format = "0"
        ws.cell(r, 5).number_format = "0.00"
        if item.startswith("F"):
            ws.cell(r, 1).fill = fill("C6EFCE")
        else:
            ws.cell(r, 1).fill = fill("BDD7EE")
    add_notes(ws, 15, 5, "Check column = Dimension / Step and should equal 0.01 mm (Focus) or 0.02 mm (Zoom).")
    ws.freeze_panes = "A4"

    # ------------------------------------------------------------------ Mechanical notes
    ws = wb.create_sheet("13 Mechanical Notes")
    set_widths(ws, [28, 22, 22, 50])
    header_bar(ws, 1, 4, "12. 机构外形图  Mechanical Outside Drawing  (numeric callouts)")
    col_headers(ws, 2, ["Item", "Value", "Tolerance", "Note"])
    mech = [
        ("Overall length", "83.88 mm", "±0.5", "PDF page 19"),
        ("Barrel diameter", "Φ40 mm", "MAX 24.9 / 24.1 noted on FPC sides", "image plane / gold finger up"),
        ("Mechanical back focus MBF", "+0.76 mm", "in air", "datum A"),
        ("Datum A runout", "0.02", "", "concentricity symbol"),
        ("Mount / body stack", "6.75 / 25.93 mm", "±0.5", ""),
        ("58.27 / 32.27 / 33.37 mm", "see drawing", "±0.2 / ±0.2 / ±0.1", "axial stations"),
        ("36.57 / 11.07 / 33.87 mm", "see drawing", "±0.2 / ±0.2 / ±0.1", ""),
        ("Holes Φ1.7 × depth 5", "multiple", "", "incl. mold glue-relief"),
        ("Unspecified linear tol.", "0–6: ±0.1; 6–30: ±0.15; 30–100: ±0.2; >100: ±0.3", "angle ±0.5°", "A2, 1:1, mm"),
        ("Drawing control", "Hunter / Skull / Owen 2021.08.24", "Rev A", "5X (104 组态) 外形图"),
    ]
    for i, row in enumerate(mech):
        apply_row(ws, 3 + i, row, fill_color=WHITE if i % 2 == 0 else ZEBRA, heights=28)
    add_notes(
        ws,
        14,
        4,
        "CAD graphics cannot be reconstructed as cells. Use the original PDF page 19 as the mechanical authority. "
        "2D viewpoint drawings are on pages 20–21; tabulated data is on sheet 14.",
    )
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------ View point
    ws = wb.create_sheet("14 View Point Data")
    set_widths(ws, [28, 16, 22, 22, 18])
    header_bar(ws, 1, 5, "13. 2D视点图  2D View Point Drawing")
    header_bar(ws, 2, 5, "5X 4K 最大视场角视点数据  View Point Data of Maximum View Angle", fill_color="2E75B6")
    col_headers(ws, 3, ["像高 Image Height", "视场角 FOV (°)", "视点高度 Height (mm)", "视点深度 Depth (mm)", "Axis"])
    max_vp = [
        ("D: φ8.81+0.4", 48.7, "φ28.45", 33.62, "DFOV 48.7°"),
        ("H: φ7.68+0.4", 43.1, "φ26.26", 35.15, "HFOV 43.1°"),
        ("V: φ4.32+0.4", 25.5, "φ20.19", 45.65, "VFOV 25.5°"),
    ]
    for i, row in enumerate(max_vp):
        apply_row(ws, 4 + i, list(row), fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=22)
        ws.cell(4 + i, 2).number_format = "0.0"
        ws.cell(4 + i, 4).number_format = "0.00"
    header_bar(ws, 8, 5, "5X 4K 最高视点视点数据  View Point Data of Highest View Point", fill_color="548235")
    col_headers(ws, 9, ["像高 Image Height", "视场角 FOV (°)", "视点高度 Height (mm)", "视点深度 Depth (mm)", "Axis"])
    high_vp = [
        ("D: φ8.81+0.4", 36.5, "φ34.17", 55.11, "DFOV 36.5°"),
        ("H: φ7.68+0.4", 17.8, "φ33.59", 110.6, "HFOV 17.8°"),
        ("V: φ4.32+0.4", 7.9, "φ32.39", 237.82, "VFOV 7.9°"),
    ]
    for i, row in enumerate(high_vp):
        apply_row(ws, 10 + i, list(row), fill_color=WHITE if i % 2 == 0 else ZEBRA, align=CENTER, heights=22)
        ws.cell(10 + i, 2).number_format = "0.0"
        ws.cell(10 + i, 4).number_format = "0.00"
    add_notes(ws, 14, 5, "Drawings dated Hunter / Skull / Owen 2021/8/25. Title on PDF uses IMAX334; sensor in optical spec is IMX334.")
    ws.freeze_panes = "A4"

    # ------------------------------------------------------------------ Iris Fno — full numeric table
    iris_rows = [
        (0, 1.29, 17.4, 237.787, "H", "L", "L", "H"),
        (1, 1.29, 17.4, 237.787, "H", "L", "H", "L"),
        (2, 1.30, 17.31, 235.302, "L", "H", "H", "L"),
        (3, 1.31, 17.21, 232.509, "L", "H", "L", "H"),
        (4, 1.31, 17.1, 229.429, "H", "L", "L", "H"),
        (5, 1.32, 16.97, 226.13, "H", "L", "H", "L"),
        (6, 1.33, 16.84, 222.644, "L", "H", "H", "L"),
        (7, 1.34, 16.7, 218.988, "L", "H", "L", "H"),
        (8, 1.36, 16.56, 215.172, "H", "L", "L", "H"),
        (9, 1.37, 16.4, 211.204, "H", "L", "H", "L"),
        (10, 1.38, 16.25, 207.189, "L", "H", "H", "L"),
        (11, 1.39, 16.09, 203.149, "L", "H", "L", "H"),
        (12, 1.41, 15.93, 199.088, "H", "L", "L", "H"),
        (13, 1.42, 15.76, 195.009, "H", "L", "H", "L"),
        (14, 1.44, 15.59, 190.912, "L", "H", "H", "L"),
        (15, 1.45, 15.43, 186.801, "L", "H", "L", "H"),
        (16, 1.47, 15.25, 182.676, "H", "L", "L", "H"),
        (17, 1.48, 15.08, 178.54, "H", "L", "H", "L"),
        (18, 1.50, 14.91, 174.395, "L", "H", "H", "L"),
        (19, 1.52, 14.73, 170.243, "L", "H", "L", "H"),
        (20, 1.54, 14.55, 166.085, "H", "L", "L", "H"),
        (21, 1.56, 14.36, 161.924, "H", "L", "H", "L"),
        (22, 1.58, 14.18, 157.762, "L", "H", "H", "L"),
        (23, 1.60, 13.99, 153.601, "L", "H", "L", "H"),
        (24, 1.62, 13.8, 149.444, "H", "L", "L", "H"),
        (25, 1.64, 13.6, 145.291, "H", "L", "H", "L"),
        (26, 1.67, 13.41, 141.147, "L", "H", "H", "L"),
        (27, 1.69, 13.21, 137.012, "L", "H", "L", "H"),
        (28, 1.72, 13.01, 132.889, "H", "L", "L", "H"),
        (29, 1.74, 12.81, 128.781, "H", "L", "H", "L"),
        (30, 1.77, 12.6, 124.689, "L", "H", "H", "L"),
        (31, 1.80, 12.4, 120.616, "L", "H", "L", "H"),
        (32, 1.83, 12.19, 116.565, "H", "L", "L", "H"),
        (33, 1.86, 11.97, 112.537, "H", "L", "H", "L"),
        (34, 1.89, 11.76, 108.536, "L", "H", "H", "L"),
        (35, 1.93, 11.54, 104.562, "L", "H", "L", "H"),
        (36, 1.97, 11.32, 100.62, "H", "L", "L", "H"),
        (37, 2.01, 11.1, 96.711, "H", "L", "H", "L"),
        (38, 2.05, 10.87, 92.838, "L", "H", "H", "L"),
        (39, 2.09, 10.65, 89.003, "L", "H", "L", "H"),
        (40, 2.13, 10.42, 85.209, "H", "L", "L", "H"),
        (41, 2.18, 10.19, 81.458, "H", "L", "H", "L"),
        (42, 2.23, 9.95, 77.753, "L", "H", "H", "L"),
        (43, 2.29, 9.72, 74.097, "L", "H", "L", "H"),
        (44, 2.34, 9.48, 70.491, "H", "L", "L", "H"),
        (45, 2.41, 9.23, 66.939, "H", "L", "H", "L"),
        (46, 2.47, 8.99, 63.444, "L", "H", "H", "L"),
        (47, 2.54, 8.74, 60.008, "L", "H", "L", "H"),
        (48, 2.61, 8.49, 56.633, "H", "L", "L", "H"),
        (49, 2.69, 8.24, 53.323, "H", "L", "H", "L"),
        (50, 2.78, 7.99, 50.081, "L", "H", "H", "L"),
        (51, 2.87, 7.73, 46.909, "L", "H", "L", "H"),
        (52, 2.97, 7.47, 43.81, "H", "L", "L", "H"),
        (53, 3.08, 7.21, 40.788, "H", "L", "H", "L"),
        (54, 3.19, 6.94, 37.845, "L", "H", "H", "L"),
        (55, 3.32, 6.68, 34.984, "L", "H", "L", "H"),
        (56, 3.46, 6.41, 32.21, "H", "L", "L", "H"),
        (57, 3.61, 6.13, 29.526, "H", "L", "H", "L"),
        (58, 3.78, 5.86, 26.935, "L", "H", "H", "L"),
        (59, 3.97, 5.58, 24.442, "L", "H", "L", "H"),
        (60, 4.18, 5.3, 22.05, "H", "L", "L", "H"),
        (61, 4.41, 5.02, 19.765, "H", "L", "H", "L"),
        (62, 4.68, 4.73, 17.591, "L", "H", "H", "L"),
        (63, 4.97, 4.45, 15.535, "L", "H", "L", "H"),
        (64, 5.32, 4.16, 13.602, "H", "L", "L", "H"),
        (65, 5.70, 3.88, 11.801, "H", "L", "H", "L"),
        (66, 6.16, 3.59, 10.136, "L", "H", "H", "L"),
        (67, 6.68, 3.31, 8.606, "L", "H", "L", "H"),
        (68, 7.30, 3.03, 7.209, "H", "L", "L", "H"),
        (69, 8.04, 2.75, 5.943, "H", "L", "H", "L"),
        (70, 8.95, 2.47, 4.806, "L", "H", "H", "L"),
        (71, 10.05, 2.2, 3.796, "L", "H", "L", "H"),
        (72, 11.46, 1.93, 2.91, "H", "L", "L", "H"),
        (73, 13.40, 1.65, 2.146, "H", "L", "H", "L"),
        (74, 16.02, 1.38, 1.503, "L", "H", "H", "L"),
        (75, 19.74, 1.12, 0.977, "L", "H", "L", "H"),
        (76, 26.01, 0.85, 0.566, "H", "L", "L", "H"),
        (77, 38.11, 0.58, 0.267, "H", "L", "H", "L"),
        (78, 69.08, 0.32, 0.078, "L", "H", "H", "L"),
        (79, 276.24, 0.08, 0.005, "L", "H", "L", "H"),
        (80, None, 0, 0, "H", "L", "L", "H"),
        (81, None, 0, 0, "H", "L", "H", "L"),
        (82, None, 0, 0, "L", "H", "H", "L"),
        (83, None, 0, 0, "L", "H", "L", "H"),
    ]

    ws = wb.create_sheet("15 Iris Fno Spec")
    set_widths(ws, [12, 10, 14, 16, 10, 10, 10, 10, 16, 18])
    header_bar(ws, 1, 10, "14. Iris & Fno 规格  (Wide)  Iris Step / Diameter / Fno")
    ws.merge_cells("A2:J2")
    ws["A2"].value = (
        "下表为 Wide 时，Iris 步数、通光孔径与 Fno 的对应关系。"
        " The following table shows the corresponding relationship between the Iris Step, Diameter and the Fno while Wide."
    )
    ws["A2"].alignment = WRAP
    ws["A2"].fill = fill(AMBER)
    ws["A2"].font = font(10, False, NAVY)
    ws.row_dimensions[2].height = 32
    col_headers(
        ws,
        3,
        ["Iris Step", "Fno", "Diameter Φ (mm)", "Open Area (mm²)", "A+", "A-", "B+", "B-", "Area check π·(Φ/2)²", "State"],
    )
    for i, (step, fno, dia, area, ap, am, bp, bm) in enumerate(iris_rows):
        r = 4 + i
        state = f"{ap}{am}{bp}{bm}"
        apply_row(
            ws,
            r,
            [step, fno, dia, area, ap, am, bp, bm, f"=PI()*(C{r}/2)^2", state],
            fill_color=WHITE if i % 2 == 0 else ZEBRA,
            align=CENTER,
            heights=18,
        )
        ws.cell(r, 1).font = font(11, True, NAVY)
        ws.cell(r, 2).number_format = "0.00"
        ws.cell(r, 3).number_format = "0.00"
        ws.cell(r, 4).number_format = "0.000"
        ws.cell(r, 9).number_format = "0.000"
        for col, val in ((5, ap), (6, am), (7, bp), (8, bm)):
            ws.cell(r, col).fill = fill(GREEN) if val == "H" else fill(AMBER)
        if fno is None:
            ws.cell(r, 2).fill = fill(RED_SOFT)
            ws.cell(r, 10).value = "closed"
    last = 4 + len(iris_rows) - 1
    ws.auto_filter.ref = f"A3:J{last}"
    ws.freeze_panes = "A4"
    ws.conditional_formatting.add(
        f"B4:B{last}",
        ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"),
    )

    chart = LineChart()
    chart.title = "Iris Step vs Fno (Wide)"
    chart.style = 10
    chart.y_axis.title = "Fno"
    chart.x_axis.title = "Iris Step"
    chart.height = 10
    chart.width = 18
    chart.y_axis.scaling.min = 0
    data = Reference(ws, min_col=2, min_row=3, max_row=83)  # steps 0–79 with Fno
    cats = Reference(ws, min_col=1, min_row=4, max_row=83)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    ws.add_chart(chart, "L3")

    chart2 = LineChart()
    chart2.title = "Iris Step vs Open Area (mm²)"
    chart2.style = 12
    chart2.y_axis.title = "Open Area (mm²)"
    chart2.x_axis.title = "Iris Step"
    chart2.height = 10
    chart2.width = 18
    data2 = Reference(ws, min_col=4, min_row=3, max_row=last)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
    ws.add_chart(chart2, "L22")

    add_notes(
        ws,
        last + 2,
        10,
        "Steps 80–83 have diameter/area = 0 (fully closed); Fno is blank in the PDF. "
        "Area check column recomputes π·(Φ/2)² from Diameter for verification against Open Area.",
    )
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"

    # ------------------------------------------------------------------ Appearance
    ws = wb.create_sheet("16 Appearance")
    set_widths(ws, [8, 28, 90])
    header_bar(ws, 1, 3, "15. 外观检验规格  Appearance Check Spec")
    col_headers(ws, 2, ["NO.", "项目 Item", "描述 Description"])
    apply_row(
        ws,
        3,
        [
            1,
            "外观检测\nAppearance Inspection",
            "1) 检测方式 Check Method: Indoor table lamp, visual. Eye-to-sample 30 cm ± 5 cm.\n"
            "2) 检查工具 Inspection tool: Scratch/dig film ruler.\n"
            "3) 检测标准 Check Standard:\n"
            "   a) 线伤/点伤 Scratches/Spots of lenses/filters: MIL 60-40. Total scratch length of a single lens or filter < 1/2 diameter. Acceptable if scratches/spots disappear after rotation.\n"
            "   b) FPC: slight creases acceptable if they do not affect lens performance.\n"
            "   c) Lens external: slight scratches acceptable if they do not affect performance.\n"
            "   d) Lens inside: refer to MIL 60-40.\n"
            "4) 参考文献 Reference: MIL-PRF-13830B performance standard.",
        ],
        fill_color=WHITE,
        heights=140,
    )
    ws.cell(3, 1).alignment = CENTER
    ws.cell(3, 1).font = font(11, True, NAVY)
    ws.freeze_panes = "A3"

    # Print defaults
    for sheet in wb.worksheets:
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddHeader.left.text = "Confidential  ·  嘉兴中润光学"
        sheet.oddHeader.right.text = "F139T-4  90.S08600.104"
        sheet.oddFooter.center.text = "&A  |  Page &P of &N"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    wb.properties.title = "F139T-4 Product Specification"
    wb.properties.subject = "1/1.8 5X Zoom/Focus Lens IMX334"
    wb.properties.creator = "PDF to Excel conversion"
    wb.properties.category = "Product Specification"
    return wb


def main():
    out_dir = Path("/workspace/docs")
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "F139T-4_1.8_5X_IMX334_Product_Specification.xlsx"
    wb = build_workbook()
    wb.save(out)
    art = Path("/opt/cursor/artifacts")
    art.mkdir(parents=True, exist_ok=True)
    art_out = art / "F139T-4_1.8_5X_IMX334_Product_Specification.xlsx"
    wb.save(art_out)
    print(f"wrote {out}")
    print(f"wrote {art_out}")
    print(f"sheets={wb.sheetnames}")


if __name__ == "__main__":
    main()
